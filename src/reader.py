"""Load a rental agreement as text or as a page image."""

import io
from dataclasses import dataclass
from pathlib import Path

from PIL import Image

from docx import Document

from src.config import DATA_DIR

TEXT_SUFFIXES = {".docx"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


@dataclass(frozen=True)
class PreparedDocument:
    file_name: str
    path: Path
    modality: str
    text: str
    image_bytes: bytes | None
    media_type: str | None


def logical_stem(path: Path) -> str:
    """CSV file name for a document, including names like ``foo.pdf.docx``."""
    name = path.name
    lower = name.lower()
    for suffix in (".pdf.docx", ".docx", ".png", ".jpg", ".jpeg", ".webp", ".pdf"):
        if lower.endswith(suffix):
            return name[: -len(suffix)]
    return path.stem


def find_document(file_name: str, data_dir: Path = DATA_DIR) -> Path:
    matches = []
    for folder in (data_dir / "train", data_dir / "test", data_dir / "try"):
        if not folder.exists():
            continue
        for path in folder.iterdir():
            if path.is_file() and logical_stem(path) == file_name:
                matches.append(path)
    if not matches:
        raise FileNotFoundError(f"No document for {file_name!r} under {data_dir}")
    return matches[0]


def _paragraphs(container) -> list[str]:
    lines = []
    for paragraph in container.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in getattr(container, "tables", []):
        for row in table.rows:
            cells = []
            seen = set()
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in seen:
                    cells.append(value)
                    seen.add(value)
            if cells:
                lines.append(" | ".join(cells))
    return lines


def read_docx(path: Path) -> str:
    document = Document(str(path))
    lines = _paragraphs(document)
    for section in document.sections:
        lines.extend(_paragraphs(section.header))
        lines.extend(_paragraphs(section.footer))
    return "\n".join(lines)


def image_for_model(data: bytes, media_type: str) -> tuple[bytes, str]:
    """Shrink a tall scan before it is sent. The model call is dominated by page size."""
    import io

    from PIL import Image

    image = Image.open(io.BytesIO(data)).convert("RGB")
    width, height = image.size
    longest = max(width, height)
    max_edge = 2000
    if longest > max_edge:
        scale = max_edge / longest
        image = image.resize(
            (max(1, int(width * scale)), max(1, int(height * scale))),
            Image.Resampling.LANCZOS,
        )
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=80, optimize=True)
    return buffer.getvalue(), "image/jpeg"


def prepare(file_name: str, data_dir: Path = DATA_DIR) -> PreparedDocument:
    path = find_document(file_name, data_dir)
    suffix = path.suffix.lower()
    if path.name.lower().endswith(".docx"):
        return PreparedDocument(
            file_name=file_name,
            path=path,
            modality="text",
            text=read_docx(path),
            image_bytes=None,
            media_type=None,
        )
    if suffix in IMAGE_SUFFIXES:
        return PreparedDocument(
            file_name=file_name,
            path=path,
            modality="image",
            text="",
            image_bytes=path.read_bytes(),
            media_type=MEDIA_TYPES[suffix],
        )
    raise ValueError(f"Unsupported file type: {path.name}")
